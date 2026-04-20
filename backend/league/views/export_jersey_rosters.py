from collections import defaultdict
import io
from django.http import HttpResponse
from django.utils.decorators import method_decorator
from django.views.decorators.csrf import csrf_exempt
from rest_framework.views import APIView
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from league.models.teams import Team
from league.models.draft_selection import DraftSelection

HEADER_FILL = PatternFill('solid', start_color='1F3864')
HEADER_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=10)
BODY_FONT   = Font(name='Arial', size=10)
CENTER      = Alignment(horizontal='center', vertical='center')
LEFT        = Alignment(horizontal='left', vertical='center')

thin = Side(style='thin', color='BFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

JERSEY_SIZES = ['YXS', 'YS', 'YM', 'YL', 'YXL', 'AS', 'AM', 'AL', 'AXL', 'AXXL']


def _style_header(ws, row, col_count):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER


def _style_body(ws, row, col_count):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = LEFT if col == 1 else CENTER


@method_decorator(csrf_exempt, name='dispatch')
class ExportJerseyRoster(APIView):
    """
    Exports a league-wide jersey/hat roster grouped by jersey color.
    Only includes teams that belong to a completed draft.
    Columns: Team Name | Coach | Division | Hat Count | YXS … AXXL
    """

    def get(self, request):
        COL_COUNT = 4 + len(JERSEY_SIZES)

        wb = Workbook()
        ws = wb.active
        ws.title = 'Jersey Roster'

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=COL_COUNT)
        title = ws.cell(row=1, column=1, value='League Jersey & Hat Roster')
        title.font = Font(name='Arial', bold=True, size=13)
        title.alignment = CENTER
        ws.row_dimensions[1].height = 24

        current_row = 2

        teams_qs = Team.objects.filter(
            drafts__is_complete=True
        ).distinct().select_related('division').order_by('jersey_color', 'name')

        color_groups = defaultdict(list)
        for team in teams_qs:
            color_key = (team.jersey_color or '—').strip()
            color_groups[color_key].append(team)

        for color, color_teams in sorted(color_groups.items()):
            # Color banner row
            ws.merge_cells(
                start_row=current_row, start_column=1,
                end_row=current_row, end_column=COL_COUNT
            )
            banner = ws.cell(row=current_row, column=1, value=f'Color: {color}')
            banner.font = Font(name='Arial', bold=True, size=11, color='FFFFFF')
            banner.fill = PatternFill('solid', start_color='2E4057')
            banner.alignment = LEFT
            ws.row_dimensions[current_row].height = 18
            current_row += 1

            # Column headers
            headers = ['Team Name', 'Coach', 'Division', 'Hat Count'] + JERSEY_SIZES
            for col, label in enumerate(headers, 1):
                ws.cell(row=current_row, column=col, value=label)
            _style_header(ws, current_row, COL_COUNT)
            ws.row_dimensions[current_row].height = 16
            current_row += 1

            group_size_totals = {s: 0 for s in JERSEY_SIZES}

            for team in color_teams:
                selections = list(
                    DraftSelection.objects.filter(
                        team=team,
                        draft__is_complete=True
                    ).select_related('player')
                )
                hat_count = len(selections)

                size_counts = {s: 0 for s in JERSEY_SIZES}
                for sel in selections:
                    size = (sel.player.jersey_size or '').upper().strip()
                    if size in size_counts:
                        size_counts[size] += 1
                        group_size_totals[size] += 1

                ws.cell(row=current_row, column=1, value=team.name)
                ws.cell(row=current_row, column=2, value=team.coach or '—')
                ws.cell(row=current_row, column=3, value=team.division.name if team.division else '—')
                ws.cell(row=current_row, column=4, value=hat_count)
                for i, size in enumerate(JERSEY_SIZES, 5):
                    count = size_counts[size]
                    ws.cell(row=current_row, column=i, value=count if count > 0 else '')
                _style_body(ws, current_row, COL_COUNT)
                current_row += 1

            # Totals row
            ws.cell(row=current_row, column=1, value='Total')
            for col in range(2, 5):
                ws.cell(row=current_row, column=col, value='')
            for i, size in enumerate(JERSEY_SIZES, 5):
                total = group_size_totals[size]
                ws.cell(row=current_row, column=i, value=total if total > 0 else '')
            for col in range(1, COL_COUNT + 1):
                cell = ws.cell(row=current_row, column=col)
                cell.font = Font(name='Arial', bold=True, size=10)
                cell.border = BORDER
                cell.alignment = LEFT if col == 1 else CENTER
            current_row += 2  # total row + spacer

        col_widths = [24, 22, 14, 10] + [7] * len(JERSEY_SIZES)
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="export_jersey_rosters.xlsx"'
        return response